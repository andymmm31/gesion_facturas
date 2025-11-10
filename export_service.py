from tkinter import filedialog
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def export_to_excel(invoice_list, summary_data):
    """
    Exporta una lista de facturas y un resumen a un archivo de Excel.
    """
    filepath = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Archivos de Excel", "*.xlsx"), ("Todos los archivos", "*.*")],
        title="Guardar reporte como..."
    )

    if not filepath:
        return False  # El usuario canceló

    try:
        wb = Workbook()

        # --- Hoja de Resumen ---
        ws_summary = wb.active
        ws_summary.title = "Resumen"

        ws_summary['A1'] = "Resumen del Reporte"
        ws_summary['A1'].font = Font(bold=True, size=14)

        row = 3
        for key, value in summary_data.items():
            ws_summary[f'A{row}'] = key
            ws_summary[f'B{row}'] = value
            if "€" in key:
                ws_summary[f'B{row}'].number_format = '#,##0.00€'
            row += 1

        # --- Hoja de Facturas ---
        ws_invoices = wb.create_sheet(title="Facturas")

        headers = ["ID", "Empresa", "Cliente", "Fecha Factura", "Base", "IVA %", "IVA", "RE %", "RE", "Factor", "Total"]
        ws_invoices.append(headers)

        for col_num, header in enumerate(headers, 1):
            cell = ws_invoices.cell(row=1, column=col_num)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for invoice in invoice_list:
            row_data = [
                invoice.get('id'),
                invoice.get('company'),
                invoice.get('client'),
                invoice.get('invoiceDate'),
                invoice.get('baseAmount'),
                invoice.get('iva_percent'),
                invoice.get('iva'),
                invoice.get('re_percent'),
                invoice.get('re'),
                invoice.get('factor'),
                invoice.get('amount')
            ]
            ws_invoices.append(row_data)

        # Autoajustar el ancho de las columnas
        for ws in [ws_summary, ws_invoices]:
            for col in ws.columns:
                max_length = 0
                column = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column].width = adjusted_width

        wb.save(filepath)
        return True
    except Exception as e:
        print(f"Error al exportar a Excel: {e}")
        return False
